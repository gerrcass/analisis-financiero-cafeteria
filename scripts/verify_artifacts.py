"""Verificador semántico de artefactos generados.

Regenera el contrato, el libro y el sitio en un directorio temporal y compara
su contenido semántico con lo commitado. No compara bytes de `.xlsx` ni de
imágenes: openpyxl y los motores de captura pueden cambiar el contenedor ZIP o
los metadatos sin cambiar el producto.

Uso:
    python scripts/verify_artifacts.py
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.build_site import build_site  # noqa: E402
from scripts.build_workbook import WORKBOOK_PATH, build_workbook  # noqa: E402
from scripts.model import EXPECTED_PATH, load_model, write_expected  # noqa: E402

BINARY_SUFFIXES = {".xlsx", ".png", ".jpg", ".jpeg", ".webp", ".gif", ".ico"}
REQUIRED_SITE_FILES = (
    "index.html",
    "aprende.html",
    "404.html",
    ".nojekyll",
    "robots.txt",
    "sitemap.xml",
    "data/model.json",
    "js/model.js",
    "js/charts.js",
    "js/app.js",
    "js/lab.js",
    "js/guide.js",
    "css/styles.css",
    "assets/aurora-mark.svg",
    "assets/icon.svg",
)


def _chart_refs(chart: Any) -> list[tuple[str, str]]:
    refs: list[tuple[str, str]] = []

    def add(container: Any, kind: str) -> None:
        for series in getattr(container, "series", []):
            value = getattr(series, "val", None)
            if value is not None and getattr(value, "numRef", None) is not None:
                refs.append((kind, value.numRef.f))
            if value is not None and getattr(value, "strRef", None) is not None:
                refs.append((kind, value.strRef.f))

    add(chart, "series")
    for subchart in getattr(chart, "_charts", []):
        add(subchart, "subchart")
    # Categorías del gráfico principal.
    for series in getattr(chart, "series", []):
        category = getattr(series, "cat", None)
        if category is not None:
            if getattr(category, "numRef", None) is not None:
                refs.append(("categories", category.numRef.f))
            if getattr(category, "strRef", None) is not None:
                refs.append(("categories", category.strRef.f))
    return refs


def workbook_signature(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    signature: dict[str, Any] = {
        "sheets": wb.sheetnames,
        "calc": (
            bool(wb.calculation.fullCalcOnLoad),
            bool(wb.calculation.forceFullCalc),
            str(wb.calculation.calcMode),
        ),
        "names": sorted((name, definition.attr_text) for name, definition in wb.defined_names.items()),
        "cells": {},
        "merges": {},
        "charts": {},
        "freeze": {},
        "print_area": {},
    }
    for ws in wb.worksheets:
        cells: dict[str, Any] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = (cell.value, cell.data_type, cell.number_format)
        signature["cells"][ws.title] = cells
        signature["merges"][ws.title] = sorted(str(value) for value in ws.merged_cells.ranges)
        signature["charts"][ws.title] = [
            (type(chart).__name__, tuple(_chart_refs(chart))) for chart in ws._charts
        ]
        signature["freeze"][ws.title] = ws.freeze_panes
        signature["print_area"][ws.title] = str(ws.print_area)
    return signature


def compare_json(reference: Path, generated: Path, problems: list[str]) -> None:
    try:
        expected = json.loads(reference.read_text(encoding="utf-8"))
        actual = json.loads(generated.read_text(encoding="utf-8"))
    except Exception as error:  # pragma: no cover - ruta de diagnóstico
        problems.append(f"No se pudo leer JSON ({reference}): {error}")
        return
    if expected != actual:
        problems.append(f"JSON distinto: {reference} frente a la regeneración")


def compare_text_tree(reference_root: Path, generated_root: Path, problems: list[str]) -> None:
    for generated in sorted(generated_root.rglob("*")):
        if not generated.is_file():
            continue
        relative = generated.relative_to(generated_root)
        target = reference_root / relative
        if generated.suffix.lower() in BINARY_SUFFIXES:
            if not target.exists():
                problems.append(f"Falta el artefacto binario: {relative}")
            continue
        if not target.exists():
            problems.append(f"Falta el archivo generado: {relative}")
            continue
        if generated.read_bytes() != target.read_bytes():
            problems.append(f"Archivo generado distinto: {relative}")

    for required in REQUIRED_SITE_FILES:
        if not (reference_root / required).exists():
            problems.append(f"Falta un archivo obligatorio del sitio: {required}")


def verify() -> list[str]:
    problems: list[str] = []
    data = load_model()
    with tempfile.TemporaryDirectory(prefix="aurora-verify-") as tmp:
        tmp_path = Path(tmp)
        generated_expected = tmp_path / "expected.json"
        write_expected(data, generated_expected)
        compare_json(EXPECTED_PATH, generated_expected, problems)

        generated_workbook = tmp_path / "workbook.xlsx"
        build_workbook(output=generated_workbook)
        if not WORKBOOK_PATH.exists():
            problems.append(f"Falta el libro commitado: {WORKBOOK_PATH}")
        else:
            try:
                if workbook_signature(WORKBOOK_PATH) != workbook_signature(generated_workbook):
                    problems.append("El libro commitado no coincide semánticamente con el regenerado")
            except Exception as error:  # pragma: no cover - diagnóstico
                problems.append(f"No se pudo comparar el libro: {error}")

        generated_site = tmp_path / "docs"
        build_site(output=generated_site)
        compare_text_tree(ROOT / "docs", generated_site, problems)

        screenshot = ROOT / "docs" / "assets" / "screenshot-dashboard.png"
        if not screenshot.exists() or screenshot.stat().st_size < 10_000:
            problems.append("Falta docs/assets/screenshot-dashboard.png o es demasiado pequeña")

        if not (ROOT / "docs" / "guia-finanzas.md").exists():
            problems.append("Falta docs/guia-finanzas.md")

    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verifica que los artefactos commitados están al día")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args(argv)
    problems = verify()
    if problems:
        for problem in problems:
            print(f"ERROR: {problem}")
        return 1
    if not args.quiet:
        print("Artefactos semánticamente actualizados: golden, libro, sitio y guía.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
