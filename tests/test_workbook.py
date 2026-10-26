"""Auditoría del libro: estructura, fórmulas vivas y recalculación.

Cita: docs/project-brief.md §8 y §11.3. El test recorre el libro generado, no
la plantilla: si una celda calculada se convierte en un valor pegado, falla.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.build_workbook import WORKBOOK_PATH, build_workbook
from scripts.recalc import compare_with_golden, errors, key_values, recalculate

ROOT = Path(__file__).resolve().parent.parent
EXPECTED = json.loads((ROOT / "data" / "expected.json").read_text(encoding="utf-8"))
SCENARIO_SHEETS = ("Pesimista", "Base", "Optimista")
MONTH_COLS = [chr(ord("C") + i) for i in range(26)] + ["AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL"]
FORMULA_FUNCTIONS = {
    "SUM",
    "SUMPRODUCT",
    "INDEX",
    "MATCH",
    "IF",
    "IFERROR",
    "MAX",
    "MIN",
    "ROUND",
    "ROUNDUP",
    "MOD",
    "DATE",
    "EDATE",
    "PMT",
    "CHOOSE",
    "TEXT",
    "ABS",
}
FORBIDDEN = ("XLOOKUP", "IFS", "LET(", "LAMBDA", "FILTER(", "SORT(", "UNIQUE(", "WEBSERVICE", "HYPERLINK")


@pytest.fixture(scope="session")
def workbook_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("workbook") / "aurora.xlsx"
    build_workbook(output=path)
    return path


@pytest.fixture(scope="session")
def wb(workbook_path):
    return load_workbook(workbook_path)


@pytest.fixture(scope="session")
def recalc(workbook_path):
    cells = recalculate(workbook_path)
    return {"cells": cells, "values": key_values(cells)}


def test_workbook_is_generated_with_the_expected_sheets(wb):
    assert wb.sheetnames == [
        "Portada",
        "Resumen",
        "Supuestos",
        "Productos",
        "Personal",
        "Costos",
        "Pesimista",
        "Base",
        "Optimista",
        "Aprende",
    ]


def test_full_recalculation_on_load(wb):
    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True


def test_scenario_selector_is_validated(wb):
    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        assert ws["B3"].value in (1, 2, 3)
        validations = list(ws.data_validations.dataValidation)
        assert any("B3" in str(validation.sqref) for validation in validations)


def test_36_month_columns_and_year_columns(wb):
    ws = wb["Base"]
    assert ws["C6"].value == "=DATE(2026,1,1)"
    assert ws["AL6"].value == "=EDATE(AK6,1)"
    assert ws["C6"].number_format == "dd/mm/yyyy"
    assert ws["AL6"].number_format == "mmm-yy"
    assert ws["AM15"].value == "=SUM(C15:N15)"
    assert ws["AN15"].value == "=SUM(O15:Z15)"
    assert ws["AO15"].value == "=SUM(AA15:AL15)"
    assert ws["AP15"].value == "=SUM(AM15:AO15)"


def test_inputs_are_confined_to_assumption_sheets(wb):
    productos_allowed = {f"{col}{row}" for col in ("A", "D", "E", "G", "H", "I") for row in range(5, 13)}
    personal_allowed = {f"{col}{row}" for col in ("B",) for row in range(5, 9)}
    personal_allowed |= {f"{col}{row}" for col in ("C", "D", "E") for row in range(5, 9)}
    personal_allowed |= {"B11", "B12"}
    costos_allowed = {f"B{row}" for row in range(6, 11)}
    costos_allowed |= {f"{col}{row}" for col in ("B", "C") for row in range(15, 21)}
    costos_allowed.add("B24")
    allowed = {
        "Productos": productos_allowed,
        "Personal": personal_allowed,
        "Costos": costos_allowed,
        "Portada": set(),
        "Resumen": {"B4"},
    }
    for sheet, coordinates in allowed.items():
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "n" and isinstance(cell.value, (int, float)):
                    if sheet == "Supuestos":
                        continue
                    assert cell.coordinate in coordinates, (
                        f"constante numérica inesperada en {sheet}!{cell.coordinate}"
                    )
        assert isinstance(coordinates, set)

    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=10, max_row=ws.max_row):
            for cell in row:
                if cell.data_type == "n" and isinstance(cell.value, (int, float)):
                    if cell.coordinate in {"B61", "B66"} and cell.value == 0:
                        continue  # saldo inicial de pérdidas: constante documentada
                    pytest.fail(f"valor pegado en {sheet}!{cell.coordinate}: {cell.value}")
        assert ws["B3"].data_type == "n"
        assert ws["C5"].data_type == "n"

def test_all_scenario_cells_are_formulas(wb):
    checked = 0
    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=12, max_row=156):
            for cell in row:
                if cell.value is None:
                    continue
                if cell.column_letter in ("A", "B"):
                    continue
                if cell.data_type == "f":
                    checked += 1
                    continue
                if cell.data_type == "n":
                    pytest.fail(f"valor pegado en {sheet}!{cell.coordinate}: {cell.value!r}")
    assert checked > 2000


def test_scenario_sheets_share_normalised_formulas(wb):
    structures = []
    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        formulas = {
            f"{cell.column_letter}{cell.row}": cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }
        structures.append(formulas)
    assert structures[0] == structures[1] == structures[2]


def test_defined_names_resolve(wb):
    assert len(wb.defined_names) >= 45
    for name, definition in wb.defined_names.items():
        target = definition.attr_text
        assert target, f"nombre vacío: {name}"
        assert "!" in target
        sheet = target.split("!")[0].strip("'")
        assert sheet in wb.sheetnames, f"{name} apunta a {sheet}"
    for required in (
        "BaseTicketsDay",
        "ScenarioMatrix",
        "LoanPayment",
        "WeightedNetPrice",
        "WeightedCostAdj",
        "StaffCostByYear",
        "MonthlyDepreciation",
        "MaintenanceCapexPct",
        "TotalCapex",
    ):
        assert required in wb.defined_names


def test_formula_subset_and_no_external_links(wb):
    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                formula = str(cell.value)
                assert "[" not in formula and ".xlsx" not in formula
                for forbidden in FORBIDDEN:
                    assert forbidden not in formula.upper(), f"{sheet}!{cell.coordinate}: {formula}"


def test_key_rows_have_the_brief_structure(wb):
    ws = wb["Base"]
    assert ws["C12"].value.startswith("=BaseTicketsDay*INDEX(ScenarioMatrix,1,$B$3)")
    assert ws["C14"].value == "=C13*ItemsPerTicket"
    assert ws["C15"].value.startswith("=C14*WeightedNetPrice")
    assert ws["C16"].value.startswith("=C14*WeightedCostAdj")
    assert ws["C17"].value == "=C15-C16"
    assert ws["C26"].value == "=C17-C25"
    assert ws["C27"].value == "=MonthlyDepreciation"
    assert ws["C28"].value == "=C26-C27"
    assert ws["C29"].value == "=C53"
    assert ws["C30"].value == "=C28-C29"
    assert ws["C31"].value == "=IF(C7=1,$B$63/12,IF(C7=2,$C$63/12,$D$63/12))"
    assert ws["C32"].value == "=C30-C31"
    assert ws["C38"].value == "=OpeningCash"
    assert ws["C39"].value == "=C16*(InventoryDays-PayableDays)/365"
    assert ws["C41"].value == "=C15*MaintenanceCapexPct"
    assert ws["C43"].value == "=C26-C42-C41-C40"
    assert ws["C47"].value == "=C43-C44+C45-C46"
    assert ws["C48"].value == "=C38+C47"
    assert ws["C52"].value == "=LoanAmount"
    assert ws["C53"].value == "=C52*LoanRate/12"
    assert ws["C54"].value == "=MIN(LoanPayment-C53,C52)"
    assert ws["C55"].value == "=C52-C54"
    assert ws["C57"].value.startswith("=IFERROR((C26-C31)/(C53+C54),0)")
    assert ws["B113"].value.startswith("=IF(MIN(C111:AL111)=999")
    assert ws["B114"].value.startswith("=IF(MIN(C112:AL112)=999")


def test_freezes_panes_and_print_setup(wb):
    for sheet in SCENARIO_SHEETS:
        ws = wb[sheet]
        assert ws.freeze_panes == "C10"
        assert ws.page_setup.orientation == "landscape"
        assert ws.print_title_rows == "$1:$9" or ws.print_title_rows == "1:9"
        assert ws.page_setup.fitToWidth == 1


def test_supuestos_contains_the_contract_checks(wb):
    ws = wb["Supuestos"]
    assert ws["B13"].value == 360
    assert ws["B14"].value == 1.35
    assert ws["B48"].value == "=PMT(B43/12,B44,-B42)"
    assert ws["B67"].value == "=SUM(Productos!H5:H12)"
    assert ws["C67"].value.startswith("=IF(ABS(B67-1)")
    assert ws["B70"].value == "=Costos!B21"
    assert ws["B71"].value == "=B41+B42"
    assert ws["B72"].value.startswith("=IF(AND(D27<E27")
    assert ws["B73"].value.startswith("=IF(AND(D30>E30")


def test_costos_sources_equal_uses(wb):
    ws = wb["Costos"]
    assert ws["B21"].value == "=SUM(B15:B20)"
    assert ws["B30"].value == "=B21+Supuestos!B46+Supuestos!B45"
    assert ws["B31"].value == "=Supuestos!B41+Supuestos!B42"
    assert ws["B32"].value.startswith("=IF(B30=B31")


def test_summary_has_kpis_comparison_and_chart_blocks(wb):
    ws = wb["Resumen"]
    assert ws["B4"].value == 2
    assert ws["C4"].value == '=CHOOSE(B4,"Pesimista","Base","Optimista")'
    assert ws["A10"].value == "Base"
    assert ws["C10"].value == "=Base!B76"
    assert ws["D10"].value == "=Base!B86"
    assert ws["E10"].value == "=IFERROR(D10/C10,0)"
    assert ws["F10"].value == "=Base!B92"
    assert ws["H10"].value == "=Base!B98"
    assert ws["I10"].value == "=Base!B103"
    assert ws["J10"].value == "=Base!B113"
    assert ws["K10"].value == "=Base!B114"
    assert ws["B18"].value == "=Pesimista!B113"
    assert ws["D18"].value == "=Pesimista!AL110"
    assert ws["B23"].value == "=Pesimista!D76"
    assert ws["C111"].value.startswith("=CHOOSE($B$4,Pesimista!C15")
    assert ws["B111"].value == "=Base!C6"
    assert ws["C151"].value == "=Pesimista!D86"
    assert ws["C158"].value == "=Base!C140"
    assert ws["C169"].value == "=Base!B77"
    assert ws["C174"].value == "=SUM(C169:C173)"


def _series_refs(chart):
    refs = []
    for series in chart.series:
        if series.val is not None and series.val.numRef is not None:
            refs.append(series.val.numRef.f)
    for subchart in getattr(chart, "_charts", []):
        for series in subchart.series:
            if series.val is not None and series.val.numRef is not None:
                refs.append(series.val.numRef.f)
    return refs


def test_five_live_charts_reference_cell_ranges(wb):
    ws = wb["Resumen"]
    assert len(ws._charts) == 5
    all_refs = []
    for chart in ws._charts:
        refs = _series_refs(chart)
        assert refs, f"gráfico sin series: {chart.title}"
        for ref in refs:
            assert "$" in ref and ":" in ref
            assert "Resumen" in ref
            all_refs.append(ref)
    joined = " ".join(all_refs)
    assert "$C$111:$C$146" in joined
    assert "$E$111:$E$146" in joined
    assert "$C$151:$C$153" in joined
    assert "$C$158:$C$165" in joined
    assert "$C$169:$C$173" in joined
    for cell in ("C111", "D111", "E111", "F111", "C151", "C158", "C169"):
        assert ws[cell].data_type == "f"


def test_aprende_sheet_has_chapters_examples_and_questions(wb):
    ws = wb["Aprende"]
    text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert "Aprende con Aurora" in text
    assert "CÓMO USAR ESTA HOJA" in text
    assert "LOS 10 CAPÍTULOS" in text
    assert "1. Qué es un modelo financiero" in text
    assert "10. Decidir con el modelo" in text
    assert "FÓRMULAS CLAVE" in text
    assert "GLOSARIO MÍNIMO" in text
    assert "ERRORES COMUNES" in text
    assert "PREGUNTAS DE LECTURA" in text
    formulas = [cell.value for row in ws.iter_rows() for cell in row if cell.data_type == "f"]
    assert len(formulas) >= 20
    assert any("TEXT(" in str(formula) for formula in formulas)
    assert any("Base!B76" in str(formula) for formula in formulas)


def test_recalc_summary_values(recalc):
    summary = recalc["values"]["summary"]
    assert summary["selector"] == 2
    assert summary["kpiBase2026"] == pytest.approx(384418.75, abs=1)


def test_recalc_has_no_error_values(recalc):
    assert errors(recalc["cells"]) == []


def test_recalc_matches_golden(recalc):
    assert compare_with_golden(recalc["values"]) == []


def test_recalc_specific_identities(recalc):
    cells = recalc["cells"]
    for sheet in ("BASE", "PESIMISTA", "OPTIMISTA"):
        values = cells[sheet]
        assert abs(float(values["N50"])) < 1e-6
        assert abs(float(values["N36"])) < 1e-6
        assert abs(float(values["C149"])) < 1e-6
        assert abs(float(values["C150"])) < 1e-6
        assert abs(float(values["C151"])) < 1e-6
        assert float(values["N48"]) > 0
        if values["B113"] != ">36":
            assert float(values["B113"]) in (34, 25)
        assert float(values["B104"]) > 0


def test_recalc_base_matches_the_brief_annual_values(recalc):
    values = recalc["values"]["base"]
    assert values["revenue2026"] == pytest.approx(384418.75, abs=1)
    assert values["ebitda2026"] == pytest.approx(61657.79, abs=1)
    assert values["closingCash2028"] == pytest.approx(132062.61, abs=1)
    assert values["paybackProject"] == 34
    assert values["paybackEquity"] == 30
    assert values["cumulativeFcfe"] == pytest.approx(110562.61, abs=1)
    assert values["dscr2028"] == pytest.approx(3.83, abs=0.01)


def test_workbook_opens_without_repairs_and_is_small(workbook_path):
    assert workbook_path.stat().st_size < 500_000
    load_workbook(workbook_path)


def test_live_formulas_change_with_an_input(tmp_path):
    """Cambiar un supuesto debe mover las cifras: no hay valores pegados."""
    from openpyxl import load_workbook as load

    source = tmp_path / "aurora.xlsx"
    build_workbook(output=source)
    wb = load(source)
    # El motor de recalculación del test se usa sobre una copia con demanda +10 %.
    wb["Supuestos"]["B13"] = 396
    modified = tmp_path / "aurora-modified.xlsx"
    wb.save(modified)
    cells = recalculate(modified)
    modified_revenue = cells["BASE"]["B76"]
    assert float(modified_revenue) == pytest.approx(384418.75 * 1.1, rel=1e-6)
